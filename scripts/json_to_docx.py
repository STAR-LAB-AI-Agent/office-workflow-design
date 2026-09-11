# -*- coding: utf-8 -*-
"""
json_to_docx.py — 把 excel_summary.py 的摘要 JSON 渲染成 Word 报告（工作流第 2 步）
用法: python scripts/json_to_docx.py --input outputs/摘要.json --output outputs/分析报告.docx
"""
import argparse
import json
import sys
from pathlib import Path

from docx import Document


def fail(msg):
    print(json.dumps({"status": "error", "error": msg}, ensure_ascii=False))
    sys.exit(1)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="把统计摘要 JSON 生成 Word 分析报告")
    parser.add_argument("--input", required=True, help="摘要 JSON 路径(excel_summary.py 的输出)")
    parser.add_argument("--output", required=True, help="输出 Word 文件路径")
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.exists():
        fail(f"输入文件不存在: {input_path}（请先运行 excel_summary.py）")

    try:
        data = json.loads(input_path.read_text(encoding="utf-8"))
    except Exception as e:
        fail(f"JSON 解析失败: {e}")

    doc = Document()
    doc.add_heading("数据分析报告", level=0)
    doc.add_paragraph(f"数据来源：{data.get('source_file', '未知')}")
    doc.add_paragraph(f"数据规模：共 {data.get('rows')} 行 × {data.get('columns')} 列")
    doc.add_paragraph("字段：" + "、".join(str(c) for c in data.get("column_names", [])))

    numeric_stats = data.get("numeric_stats", {})
    if numeric_stats:
        doc.add_heading("一、数值指标统计", level=1)
        table = doc.add_table(rows=1, cols=5)
        table.style = "Light Grid Accent 1"
        for i, h in enumerate(["指标", "合计", "平均值", "最大值", "最小值"]):
            table.rows[0].cells[i].text = h
        for col, st in numeric_stats.items():
            cells = table.add_row().cells
            cells[0].text = str(col)
            cells[1].text = str(st.get("sum"))
            cells[2].text = str(st.get("mean"))
            cells[3].text = str(st.get("max"))
            cells[4].text = str(st.get("min"))

    group_stats = data.get("group_stats", {})
    for idx, (col, items) in enumerate(group_stats.items(), start=2):
        doc.add_heading(f"{idx}、按「{col}」汇总", level=1)
        table = doc.add_table(rows=1, cols=2)
        table.style = "Light Grid Accent 1"
        table.rows[0].cells[0].text = str(col)
        table.rows[0].cells[1].text = "数值"
        for item in items:
            cells = table.add_row().cells
            cells[0].text = str(item["name"])
            cells[1].text = str(item["value"])

    doc.add_paragraph("")
    doc.add_paragraph("（本报告由 AI 办公流程自动生成）")

    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)

    print(json.dumps({
        "status": "success",
        "input": str(input_path),
        "outputs": {"docx": str(output_path)},
        "summary": f"Word 报告已生成：{output_path}（含数值统计表和 {len(group_stats)} 组分组汇总）",
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()