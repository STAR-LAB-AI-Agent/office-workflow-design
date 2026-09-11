# -*- coding: utf-8 -*-
"""
json_to_pptx.py — 把摘要 JSON 渲染成汇报 PPT（工作流第 3 步）
用法: python scripts/json_to_pptx.py --input outputs/摘要.json --output outputs/汇报.pptx
"""
import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt


def fail(msg):
    print(json.dumps({"status": "error", "error": msg}, ensure_ascii=False))
    sys.exit(1)


def add_bullet_slide(prs, title, lines):
    """加一页：标题 + 要点列表"""
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    tf = slide.placeholders[1].text_frame
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(18)
    return slide


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="把统计摘要 JSON 生成汇报 PPT")
    parser.add_argument("--input", required=True, help="摘要 JSON 路径(excel_summary.py 的输出)")
    parser.add_argument("--output", required=True, help="输出 PPT 文件路径")
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.exists():
        fail(f"输入文件不存在: {input_path}（请先运行 excel_summary.py）")

    try:
        data = json.loads(input_path.read_text(encoding="utf-8"))
    except Exception as e:
        fail(f"JSON 解析失败: {e}")

    prs = Presentation()

    # 封面页
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "数据分析汇报"
    slide.placeholders[1].text = (
        f"数据来源：{data.get('source_file', '未知')}\n"
        f"共 {data.get('rows')} 行 × {data.get('columns')} 列\n"
        "由 AI 办公流程自动生成"
    )

    # 数值统计页
    numeric_stats = data.get("numeric_stats", {})
    if numeric_stats:
        lines = []
        for col, st in numeric_stats.items():
            lines.append(f"{col}：合计 {st['sum']}，平均 {st['mean']}，"
                         f"最大 {st['max']}，最小 {st['min']}")
        add_bullet_slide(prs, "数值指标统计", lines)

    # 每个分组汇总一页
    for col, items in data.get("group_stats", {}).items():
        lines = [f"{it['name']}：{it['value']}" for it in items]
        add_bullet_slide(prs, f"按「{col}」汇总", lines)

    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)

    print(json.dumps({
        "status": "success",
        "input": str(input_path),
        "outputs": {"pptx": str(output_path)},
        "summary": f"PPT 已生成：{output_path}（共 {len(prs.slides)} 页）",
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()